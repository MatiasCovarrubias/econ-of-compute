#!/usr/bin/env python3
"""Build the Stage 1 39-sector VLW object: DC_facilities and Compute.

Reads one annual slice of the 37-sector network, the C30 data-center
construction series, and BEA FA detail. Writes
data/derived/calibration_data_dc.mat.

Run from anywhere:
    python3 code/build_stage1.py [--year 2018]

The 37-sector base currently exists only for 1947-2018 (VLW's
calibration_data.mat). Years after 2018 need the rebuilt base network
(stage1_plan.md S1) and stop with a message until it exists.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import numpy as np

sys.path.insert(0, str(Path(__file__).resolve().parent))
from cost_buildup import node_costs
from facilities_stock import c30_annual_m, facility_stock, office_price_index
from it_stock import it_stock
from stage1_assumptions import (
    ADMIN,
    BETA,
    C30_SCALE_RENT,
    COLO_RENT_2017_M,
    COMP,
    COMPUTE_OTHER_M_INFO_SHARE,
    COMPUTER_ASSET_CODES,
    COMPUTERS_PRODUCERS_SHARE,
    CONSTR,
    CONS_COMPUTE_SHARE,
    DC,
    DELTA_IT,
    DELTA_PLANT,
    DELTA_SHELL,
    DEPIND,
    INFO,
    IT_MARGIN_RT,
    IT_MARGIN_TRANSP,
    IT_MARGIN_WT,
    LABS_PT_SHARE,
    N37,
    N39,
    OWN_USE_INFO_SHARE,
    PT,
    R,
    RE,
    RT,
    SECTOR_NAMES_39,
    SERVER_ASSET_CODES,
    SOFT_TO_IT_RATIO,
    SOFTWARE_ASSET_CODES,
    STRUCTURES_FROM_INFO_SHARE,
    TRANSP,
    UTIL,
    WT,
    CMP,
)

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
VLW_RAW = ROOT / "data" / "raw" / "vlw"  # detailnonres_inv1.xlsx from the VLW replication packet
OUT_DIR = ROOT / "data" / "derived"

YEAR0 = 1947
BASE_LAST_YEAR = 2018  # last year in VLW's calibration_data.mat
YEAR = 2018
YIDX = YEAR - YEAR0


def set_year(year: int) -> None:
    global YEAR, YIDX
    YEAR = year
    YIDX = year - YEAR0


def _need_scipy():
    try:
        import scipy.io as sio  # noqa: F401
    except ImportError as exc:
        raise SystemExit(
            "scipy is required to read/write .mat files "
            "(pip install scipy, or PYTHONPATH=/tmp/pydeps)."
        ) from exc
    return sio


def load_2018_slice(mat_path: Path) -> dict:
    sio = _need_scipy()
    d = sio.loadmat(mat_path, squeeze_me=True)
    go = np.asarray(d["GO47bea"][YIDX], dtype=float)
    va = np.asarray(d["VA47bea"][YIDX], dtype=float)
    cons = np.asarray(d["Cons47bea"][YIDX], dtype=float)
    iomat = np.asarray(d["IOmat47bea"][:, :, YIDX], dtype=float)
    invmat = np.asarray(d["invmat"][:, :, YIDX], dtype=float)
    deprat = np.asarray(d["depratbea"][:, YIDX], dtype=float)
    capsh_all = np.asarray(d["capshbea"], dtype=float)
    # capshbea is 37 x 71 (1947–2017). Use 2017 for the 2018 slice.
    capsh = capsh_all[:, -1].copy()
    residual = go - va - iomat.sum(axis=0)
    return {
        "GO": go,
        "VA": va,
        "Cons": cons,
        "IOmat": iomat,
        "invmat": invmat,
        "deprat": deprat,
        "capsh": capsh,
        "residual": residual,
        "capsh_year": 2017,
    }


def _fa_year_col(header_row, year: int) -> int:
    for j, v in enumerate(header_row):
        if v == year or v == str(year):
            return j
        try:
            if float(v) == float(year):
                return j
        except (TypeError, ValueError):
            pass
    start = header_row[2]
    try:
        if int(float(start)) == 1901:
            return 2 + (year - 1901)
    except (TypeError, ValueError):
        pass
    raise KeyError(year)


def load_fa_year(path: Path, year: int, codes: tuple[str, ...]) -> dict[str, dict[str, float]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = [s for s in wb.sheetnames if s not in ("readme", "Datasets")]
    out = {s: {c: 0.0 for c in codes} for s in sheets}
    for s in sheets:
        rows = list(wb[s].iter_rows(values_only=True))
        j = _fa_year_col(rows[5], year)
        for r in rows[7:]:
            if r[0] is None:
                continue
            code = str(r[0]).strip()
            if code in out[s]:
                try:
                    out[s][code] = float(r[j] or 0.0)
                except (TypeError, ValueError):
                    out[s][code] = 0.0
    wb.close()
    return out


def fa_index_to_vlw() -> np.ndarray:
    """Map 0-based FA sheet index (0..62) to 0-based VLW sector, or -1 if dropped."""
    m = np.full(63, -1, dtype=int)
    for sec, spec in enumerate(DEPIND):
        if isinstance(spec, tuple):
            lo, hi = spec
            for fa in range(lo, hi + 1):
                m[fa - 1] = sec
        else:
            m[spec - 1] = sec
    return m


def computers_software_by_vlw(fa_inv: dict[str, dict[str, float]]) -> np.ndarray:
    sheets = list(fa_inv)
    mapped = fa_index_to_vlw()
    w = np.zeros(N37)
    for i, s in enumerate(sheets):
        sec = mapped[i]
        if sec < 0:
            continue
        row = fa_inv[s]
        w[sec] += sum(row[c] for c in COMPUTER_ASSET_CODES) + sum(
            row[c] for c in SOFTWARE_ASSET_CODES
        )
    return w


def take(cell: float, amount: float) -> tuple[float, float, float]:
    """Return (cell leftover, taken, unmet). Never negative."""
    taken = min(max(cell, 0.0), max(amount, 0.0))
    return cell - taken, taken, amount - taken


def i_ss(va: float, capsh: float, delta: float, r: float = R) -> float:
    k = capsh * va / (r + delta)
    return delta * k


def build_node(fac: dict[int, dict[str, float]], it: dict[int, dict[str, float]]) -> dict:
    if YEAR not in it:
        raise RuntimeError(f"IT stock not available for {YEAR} (FA detail ends 2024)")
    f = fac[YEAR]
    k_shell, k_plant = f["K_shell"], f["K_plant"]
    k_fac = k_shell + k_plant
    k_it = it[YEAR]["K_IT"]
    k_soft = it[YEAR]["K_soft"]

    costs = node_costs(YEAR, f, it[YEAR])
    go_c = costs["GO_cmp"]

    return {
        "K_shell": k_shell,
        "K_plant": k_plant,
        "K_fac": k_fac,
        "K_IT": k_it,
        "K_IT_ceiling": it[YEAR]["K_IT_ceiling"],
        "K_soft": k_soft,
        "K_fac_path": {str(y): v["K_fac"] for y, v in fac.items()},
        "K_fac_anchor_share": f["anchor_share"],
        "I_C30": f["I"],
        "TWh_all": costs["TWh_all"],
        "elec_share": costs["share"],
        "elec_price": costs["price"],
        "GOS_fac": costs["GOS_fac"],
        "elec": costs["elec"],
        "other_M_fac": costs["other_M_fac"],
        "L_fac": costs["L_fac"],
        "GO_fac": costs["GO_fac"],
        "VA_fac": costs["VA_fac"],
        "GOS_cmp": costs["GOS_cmp"],
        "other_M_cmp": costs["other_M_cmp"],
        "L_cmp": costs["L_cmp"],
        "GO_cmp": go_c,
        "VA_cmp": costs["VA_cmp"],
        "I_fac": costs["I_fac"],
        "I_IT": costs["I_IT"],
        "I_soft": costs["I_soft"],
        "delta_fac": costs["delta_fac"],
        "delta_cmp": costs["delta_cmp"],
        "capsh_fac": costs["capsh_fac"],
        "capsh_cmp": costs["capsh_cmp"],
        "own_use": OWN_USE_INFO_SHARE * go_c,
        "labs": LABS_PT_SHARE * go_c,
        "cons_cmp": CONS_COMPUTE_SHARE * go_c,
        "market": (1.0 - OWN_USE_INFO_SHARE) * go_c,
        "colo_rent": COLO_RENT_2017_M * C30_SCALE_RENT,
    }


def expand37(slice37: dict) -> dict:
    iomat = np.zeros((N39, N39))
    invmat = np.zeros((N39, N39))
    iomat[:N37, :N37] = slice37["IOmat"]
    invmat[:N37, :N37] = slice37["invmat"]
    go = np.zeros(N39)
    va = np.zeros(N39)
    cons = np.zeros(N39)
    capsh = np.zeros(N39)
    deprat = np.zeros(N39)
    residual = np.zeros(N39)
    go[:N37] = slice37["GO"]
    va[:N37] = slice37["VA"]
    cons[:N37] = slice37["Cons"]
    capsh[:N37] = slice37["capsh"]
    deprat[:N37] = slice37["deprat"]
    residual[:N37] = slice37["residual"]
    return {
        "IOmat": iomat,
        "invmat": invmat,
        "GO": go,
        "VA": va,
        "Cons": cons,
        "capsh": capsh,
        "deprat": deprat,
        "residual": residual,
    }


def place_node_columns(t: dict, node: dict) -> None:
    io, go_f, go_c = t["IOmat"], node["GO_fac"], node["GO_cmp"]
    io[UTIL, DC] = node["elec"]
    water = 0.30 * node["other_M_fac"]
    io[UTIL, DC] += water
    io[PT, DC] = 0.40 * node["other_M_fac"]
    io[ADMIN, DC] = 0.30 * node["other_M_fac"]

    io[DC, CMP] = go_f
    io[INFO, CMP] = COMPUTE_OTHER_M_INFO_SHARE * node["other_M_cmp"]
    io[PT, CMP] = (1.0 - COMPUTE_OTHER_M_INFO_SHARE) * node["other_M_cmp"]

    t["GO"][DC] = go_f
    t["VA"][DC] = node["VA_fac"]
    t["GO"][CMP] = go_c
    t["VA"][CMP] = node["VA_cmp"]
    t["Cons"][DC] = 0.0
    t["Cons"][CMP] = node["cons_cmp"]
    t["capsh"][DC] = node["capsh_fac"]
    t["capsh"][CMP] = node["capsh_cmp"]
    t["deprat"][DC] = node["delta_fac"]
    t["deprat"][CMP] = node["delta_cmp"]
    t["residual"][DC] = 0.0
    t["residual"][CMP] = 0.0


def place_compute_row(t: dict, node: dict, fa_weights: np.ndarray) -> np.ndarray:
    """Fill Compute sales. Returns the intended industry sales vector (37)."""
    go_c = node["GO_cmp"]
    sales = np.zeros(N37)
    sales[INFO] = node["own_use"]
    sales[PT] = node["labs"]
    rest = go_c - node["own_use"] - node["labs"] - node["cons_cmp"]
    w = fa_weights.copy()
    w[INFO] = 0.0
    w[PT] = 0.0
    if w.sum() <= 0:
        raise RuntimeError("FA computers+software weights are empty outside Info/Prof_Tech")
    w = w / w.sum()
    sales += rest * w
    t["IOmat"][CMP, :N37] = sales
    return sales


def place_invmat_node(t: dict, node: dict) -> None:
    # Facility investment is construction put in place (shell and installed
    # plant): all of it is bought from Construction, as for any nonresidential
    # structure in the VLW investment network.
    t["invmat"][:, DC] = 0.0
    t["invmat"][CONSTR, DC] = 1.0

    i_c = node["I_IT"] + node["I_soft"]
    t["invmat"][:, CMP] = 0.0
    t["invmat"][COMP, CMP] = COMPUTERS_PRODUCERS_SHARE * node["I_IT"] / i_c
    t["invmat"][WT, CMP] = IT_MARGIN_WT * node["I_IT"] / i_c
    t["invmat"][TRANSP, CMP] = IT_MARGIN_TRANSP * node["I_IT"] / i_c
    t["invmat"][RT, CMP] = IT_MARGIN_RT * node["I_IT"] / i_c
    t["invmat"][PT, CMP] = node["I_soft"] / i_c


def carve(t: dict, node: dict, slice37: dict) -> dict:
    io = t["IOmat"]
    report = {}

    # ORE-first: the whole facilities utilities bill (power + water).
    util_need = float(io[UTIL, DC])
    io[UTIL, RE], from_re, util_need = take(io[UTIL, RE], util_need)
    io[UTIL, INFO], from_info, util_need = take(io[UTIL, INFO], util_need)
    report["elec_from_RE"] = from_re
    report["elec_from_Info"] = from_info
    report["elec_unmet"] = util_need

    rent_need = min(node["colo_rent"], node["GO_fac"])
    io[RE, INFO], rent_taken, rent_unmet = take(io[RE, INFO], rent_need)
    report["rent_from_RE"] = rent_taken
    report["rent_unmet"] = rent_unmet

    sales = io[CMP, :N37].copy()
    replaced = np.zeros(N37)
    unmet_row = np.zeros(N37)
    for j in range(N37):
        if j == INFO:
            continue  # own-use is new M, not a reclassified Info sale
        io[INFO, j], taken, unmet = take(io[INFO, j], sales[j])
        replaced[j] = taken
        unmet_row[j] = unmet
    cons_need = node["cons_cmp"]
    t["Cons"][INFO], cons_taken, cons_unmet = take(t["Cons"][INFO], cons_need)
    report["info_row_replaced"] = float(replaced.sum())
    report["info_row_unmet"] = float(unmet_row.sum())
    report["cons_from_Info"] = cons_taken
    report["cons_unmet"] = cons_unmet

    y = node["market"]
    x = node["own_use"]
    go0 = slice37["GO"][INFO]
    va0 = slice37["VA"][INFO]
    col0 = slice37["IOmat"][:, INFO]
    scale = max(go0 - y, 0.0) / go0
    io[:N37, INFO] = col0 * scale
    io[UTIL, INFO] = col0[UTIL] - from_info
    io[CMP, INFO] = x
    t["GO"][INFO] = go0 - y
    report["info_scale"] = scale
    report["info_GO_cut"] = y
    report["info_own_use"] = x

    t["GO"][RE] = slice37["GO"][RE] - rent_taken - from_re
    va_re0 = slice37["VA"][RE]
    go_re0 = slice37["GO"][RE]
    va_re_cut = rent_taken * (va_re0 / go_re0)
    t["VA"][RE] = va_re0 - va_re_cut
    va_node = node["VA_fac"] + node["VA_cmp"]
    t["VA"][INFO] = va0 - (va_node - va_re_cut)
    report["VA_RE_cut"] = va_re_cut
    report["VA_Info_cut"] = va_node - va_re_cut

    for j in range(N37):
        extra = unmet_row[j]
        if extra > 0 and j != INFO:
            t["VA"][j] = t["VA"][j] - extra
            t["GO"][j] = slice37["GO"][j]

    i_info0 = i_ss(slice37["VA"][INFO], slice37["capsh"][INFO], slice37["deprat"][INFO])
    i_re0 = i_ss(slice37["VA"][RE], slice37["capsh"][RE], slice37["deprat"][RE])
    flow_info = slice37["invmat"][:, INFO] * i_info0
    flow_re = slice37["invmat"][:, RE] * i_re0
    flow_info[CONSTR] = max(0.0, flow_info[CONSTR] - STRUCTURES_FROM_INFO_SHARE * node["I_fac"])
    flow_re[CONSTR] = max(0.0, flow_re[CONSTR] - (1.0 - STRUCTURES_FROM_INFO_SHARE) * node["I_fac"])
    flow_info[COMP] = max(0.0, flow_info[COMP] - COMPUTERS_PRODUCERS_SHARE * node["I_IT"])
    flow_info[WT] = max(0.0, flow_info[WT] - IT_MARGIN_WT * node["I_IT"])
    flow_info[TRANSP] = max(0.0, flow_info[TRANSP] - IT_MARGIN_TRANSP * node["I_IT"])
    flow_info[RT] = max(0.0, flow_info[RT] - IT_MARGIN_RT * node["I_IT"])
    flow_info[PT] = max(0.0, flow_info[PT] - node["I_soft"])

    t["invmat"][:N37, INFO] = flow_info / flow_info.sum()
    t["invmat"][:N37, RE] = flow_re / flow_re.sum()
    report["I_ss_Info_old"] = i_info0
    report["I_ss_Info_new"] = float(flow_info.sum())
    report["I_ss_RE_old"] = i_re0
    report["I_ss_RE_new"] = float(flow_re.sum())

    k_info = slice37["capsh"][INFO] * slice37["VA"][INFO] / (R + slice37["deprat"][INFO])
    k_leave = node["K_IT"] + STRUCTURES_FROM_INFO_SHARE * node["K_fac"]
    k_stay = max(k_info - k_leave, 0.05 * k_info)
    t["deprat"][INFO] = (
        slice37["deprat"][INFO] * k_info
        - DELTA_IT * node["K_IT"]
        - DELTA_SHELL * STRUCTURES_FROM_INFO_SHARE * node["K_shell"]
        - DELTA_PLANT * STRUCTURES_FROM_INFO_SHARE * node["K_plant"]
    ) / k_stay
    t["deprat"][INFO] = float(np.clip(t["deprat"][INFO], 0.02, 0.25))
    t["capsh"][INFO] = float(np.clip(slice37["capsh"][INFO] * k_stay / k_info, 0.05, 0.95))
    report["K_info_proxy_old"] = k_info
    report["K_info_leave"] = k_leave

    t["GO"][DC] = node["GO_fac"]
    t["GO"][CMP] = node["GO_cmp"]
    va_gap = float(t["VA"].sum() - slice37["VA"].sum())
    t["VA"][INFO] -= va_gap
    report["VA_closer"] = va_gap
    return report


def identities(t: dict, slice37: dict, node: dict, carve_report: dict) -> dict:
    io, inv = t["IOmat"], t["invmat"]
    checks = {}
    checks["invmat_col_sums"] = inv.sum(axis=0).tolist()
    checks["invmat_col_ok"] = bool(np.allclose(inv.sum(axis=0), 1.0, atol=1e-9))
    checks["nonneg_IOmat"] = bool(np.all(io >= -1e-9))
    checks["nonneg_invmat"] = bool(np.all(inv >= -1e-9))
    checks["nonneg_GO_VA_Cons"] = bool(
        np.all(t["GO"] >= -1e-6) and np.all(t["VA"] >= -1e-6) and np.all(t["Cons"] >= -1e-6)
    )
    checks["Cons_DC"] = float(t["Cons"][DC])
    checks["DC_row_only_Compute"] = bool(
        abs(io[DC, CMP] - node["GO_fac"]) < 1e-4 and abs(io[DC].sum() - node["GO_fac"]) < 1e-4
    )
    util_old = float(slice37["IOmat"][UTIL].sum())
    util_new = float(io[UTIL].sum())
    checks["utilities_row_old"] = util_old
    checks["utilities_row_new"] = util_new
    checks["utilities_row_gap"] = util_new - util_old
    va_old = float(slice37["VA"].sum())
    va_new = float(t["VA"].sum())
    checks["VA_old"] = va_old
    checks["VA_new"] = va_new
    checks["VA_gap"] = va_new - va_old
    go_old = float(slice37["GO"].sum())
    go_new = float(t["GO"].sum())
    checks["GO_old"] = go_old
    checks["GO_new"] = go_new
    checks["GO_gap"] = go_new - go_old
    checks["GO_gap_vs_x_plus_DC"] = (go_new - go_old) - (node["own_use"] + node["GO_fac"])
    checks["n_sectors"] = int(t["GO"].size)
    checks["Info_VA"] = float(t["VA"][INFO])
    checks["RE_VA"] = float(t["VA"][RE])
    checks["Compute_row_plus_Cons"] = float(io[CMP].sum() + t["Cons"][CMP])
    checks["Compute_GO"] = float(t["GO"][CMP])
    checks.update({f"carve_{k}": v for k, v in carve_report.items() if np.isscalar(v)})
    return checks


def audit_cells(slice37: dict) -> dict:
    io = slice37["IOmat"]
    return {
        "Utilities_to_Info": float(io[UTIL, INFO]),
        "Utilities_to_RE": float(io[UTIL, RE]),
        "RE_to_Info": float(io[RE, INFO]),
        "Info_GO": float(slice37["GO"][INFO]),
        "Info_VA": float(slice37["VA"][INFO]),
        "Info_Cons": float(slice37["Cons"][INFO]),
        "RE_GO": float(slice37["GO"][RE]),
        "invmat_col_sums_minmax": [
            float(slice37["invmat"].sum(0).min()),
            float(slice37["invmat"].sum(0).max()),
        ],
    }


def write_mat(path: Path, t: dict, slice37: dict, node: dict, checks: dict) -> None:
    sio = _need_scipy()
    names = np.array(SECTOR_NAMES_39, dtype=object)
    payload = {
        "IOmat": t["IOmat"],
        "invmat": t["invmat"],
        "GO": t["GO"].reshape(-1, 1),
        "VA": t["VA"].reshape(-1, 1),
        "Cons": t["Cons"].reshape(-1, 1),
        "capsh": t["capsh"].reshape(-1, 1),
        "deprat": t["deprat"].reshape(-1, 1),
        "IOmat37": slice37["IOmat"],
        "invmat37": slice37["invmat"],
        "GO37": slice37["GO"].reshape(-1, 1),
        "VA37": slice37["VA"].reshape(-1, 1),
        "Cons37": slice37["Cons"].reshape(-1, 1),
        "sector_names": names,
        "year": np.array([[YEAR]], dtype=np.int32),
        "n_sectors": np.array([[N39]], dtype=np.int32),
        "beta": np.array([[BETA]]),
        "r": np.array([[R]]),
        "D3b_owned_campuses_as_rentals": np.array([[1]], dtype=np.int32),
    }
    sio.savemat(path, payload, do_compression=True)
    report = {
        "year": YEAR,
        "units": "millions of current dollars",
        "D3b": "owned hyperscale campuses rented through DC_facilities",
        "node": {k: (v if np.isscalar(v) else v) for k, v in node.items() if k != "K_fac_path"},
        "K_fac_path": node["K_fac_path"],
        "audit_2018_slice": audit_cells(slice37),
        "identities": {
            k: v
            for k, v in checks.items()
            if k != "invmat_col_sums"
        },
        "invmat_DC": t["invmat"][:, DC].tolist(),
        "invmat_Compute": t["invmat"][:, CMP].tolist(),
        "GO_new_sectors": {"DC_facilities": float(t["GO"][DC]), "Compute": float(t["GO"][CMP])},
        "VA_new_sectors": {"DC_facilities": float(t["VA"][DC]), "Compute": float(t["VA"][CMP])},
    }
    (path.with_suffix(".json")).write_text(json.dumps(report, indent=2, default=float))


def _fmt_bn(x: float) -> str:
    return f"{x / 1000.0:8.2f}"


def print_summary(t: dict, node: dict, checks: dict, slice37: dict) -> None:
    print(f"\nStage 1 {YEAR} node (billions of current dollars)")
    print(f"  r = {R:.6f}  (beta={BETA})")
    print(f"  K_fac {node['K_fac']/1000:.2f} (shell {node['K_shell']/1000:.2f} + plant {node['K_plant']/1000:.2f}, delta {node['delta_fac']:.3f}, pre-2014 anchor {node['K_fac_anchor_share']:.0%})")
    print(f"  K_IT {node['K_IT']/1000:.2f} (5110+5140 ceiling {node['K_IT_ceiling']/1000:.2f})   K_soft {node['K_soft']/1000:.2f} @ {SOFT_TO_IT_RATIO:.2f}")
    print(f"  elec {node['TWh_all']:.0f} TWh × {node['elec_share']:.2f} × {node['elec_price']*100:.1f}¢ → {_fmt_bn(node['elec'])} bn")
    print(f"  DC_facilities  GO {_fmt_bn(node['GO_fac'])}  VA {_fmt_bn(node['VA_fac'])}  elec {_fmt_bn(node['elec'])}  I_ss {_fmt_bn(node['I_fac'])}")
    print(f"  Compute        GO {_fmt_bn(node['GO_cmp'])}  VA {_fmt_bn(node['VA_cmp'])}  own-use {_fmt_bn(node['own_use'])}  I_ss {_fmt_bn(node['I_IT']+node['I_soft'])}")
    print("\nIdentities")
    print(f"  invmat columns sum to 1: {checks['invmat_col_ok']}")
    print(f"  no negatives: IO {checks['nonneg_IOmat']}  inv {checks['nonneg_invmat']}  GO/VA/C {checks['nonneg_GO_VA_Cons']}")
    print(f"  DC row = GO sold to Compute: {checks['DC_row_only_Compute']}")
    print(f"  Utilities row gap (want ~0): {checks['utilities_row_gap']/1000:.3f} bn")
    print(f"  ΣVA gap: {checks['VA_gap']/1000:.2f} bn   (old {checks['VA_old']/1000:.1f})")
    print(f"  ΣGO gap: {checks['GO_gap']/1000:.2f} bn   vs own-use+DC { (node['own_use']+node['GO_fac'])/1000:.2f}")
    print(f"  Compute row+Cons vs GO: {checks['Compute_row_plus_Cons']/1000:.3f} vs {checks['Compute_GO']/1000:.3f}")
    print(f"  Info VA {t['VA'][INFO]/1000:.1f} (was {slice37['VA'][INFO]/1000:.1f})   RE VA {t['VA'][RE]/1000:.1f} (was {slice37['VA'][RE]/1000:.1f})")
    print(f"  C30 {YEAR} construction {node['I_C30']/1000:.2f}  vs δK_fac {node['I_fac']/1000:.2f}  ratio {node['I_C30']/node['I_fac']:.2f}")
    print(f"  elec from RE {checks['carve_elec_from_RE']/1000:.2f}  Info {checks['carve_elec_from_Info']/1000:.2f}  unmet {checks['carve_elec_unmet']/1000:.3f}")
    print(f"  colo rent taken {checks['carve_rent_from_RE']/1000:.2f}")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--year", type=int, default=2018, help="base year of the 37-sector slice")
    set_year(ap.parse_args().year)
    if YEAR > BASE_LAST_YEAR:
        raise SystemExit(
            f"{YEAR}: the 37-sector base network exists only through {BASE_LAST_YEAR} "
            "(VLW calibration_data.mat). Rebuild the base from data/raw/bea_2024 first "
            "(stage1_plan.md S1). facilities_stock.py covers 2014-2025; it_stock.py covers 2018-2024."
        )

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    mat_in = RAW / "calibration_data.mat"
    fa_inv_path = VLW_RAW / "detailnonres_inv1.xlsx"
    for p in (mat_in, fa_inv_path):
        if not p.exists():
            raise SystemExit(f"missing {p}")

    slice37 = load_2018_slice(mat_in)
    audit = audit_cells(slice37)
    print(f"S1 {YEAR} slice (millions):", json.dumps(audit, indent=2))
    if YEAR == 2018 and abs(audit["Utilities_to_Info"] - 4506) > 5:
        print("warning: Utilities→Info is not the audited 4506")

    fac = facility_stock(c30_annual_m(), office_price_index())
    print("Facility stock K_fac $bn:", {y: round(v["K_fac"] / 1000, 1) for y, v in fac.items()})
    it = it_stock()
    print("IT stock K_IT $bn (514 servers):", {y: round(v["K_IT"] / 1000, 1) for y, v in it.items()})

    codes = COMPUTER_ASSET_CODES + SOFTWARE_ASSET_CODES + SERVER_ASSET_CODES
    print(f"Loading FA {YEAR} inv (Compute sales weights, VLW vintage)...")
    fa_inv = load_fa_year(fa_inv_path, YEAR, codes)
    fa_w = computers_software_by_vlw(fa_inv)

    node = build_node(fac, it)
    t = expand37(slice37)
    place_node_columns(t, node)
    place_compute_row(t, node, fa_w)
    place_invmat_node(t, node)
    carve_report = carve(t, node, slice37)
    checks = identities(t, slice37, node, carve_report)
    print_summary(t, node, checks, slice37)

    out = OUT_DIR / "calibration_data_dc.mat"
    write_mat(out, t, slice37, node, checks)
    print(f"\nWrote {out} and {out.with_suffix('.json')}")

    failed = []
    if not checks["invmat_col_ok"]:
        failed.append("invmat columns")
    if not checks["nonneg_IOmat"] or not checks["nonneg_invmat"] or not checks["nonneg_GO_VA_Cons"]:
        failed.append("negatives")
    if not checks["DC_row_only_Compute"]:
        failed.append("DC row")
    if abs(checks["utilities_row_gap"]) > 1.0:
        failed.append("utilities row")
    if abs(checks["VA_gap"]) > 1.0:
        failed.append("ΣVA")
    if t["VA"][INFO] <= 0 or t["VA"][RE] <= 0:
        failed.append("Info/RE VA")
    if abs(checks["Compute_row_plus_Cons"] - checks["Compute_GO"]) > 1.0:
        failed.append("Compute sales vs GO")
    if failed:
        print("FAILED:", ", ".join(failed))
        return 1
    print("Stage 1 tables: identities held.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
