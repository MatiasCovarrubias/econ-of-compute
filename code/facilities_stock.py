#!/usr/bin/env python3
"""Stock of data-center facilities (shell + installed power/cooling plant).

Source flow: Census C30 private construction put in place, "Data center" line,
monthly SAAR from January 2014 (data/raw/privsatime.xlsx). Construction cost
includes the mechanical and electrical plant and excludes racks and servers,
so the cumulated stock is a *facility* stock, not a BEA-style shell.

Method: deflate the annual flow with the NIPA office-structures price index,
split each year's flow into shell and plant with a fixed cost share, run a
perpetual inventory on each part with its own depreciation rate, and revalue
to current cost. The end-2013 stock is I_2014/(delta+g) with a pre-sample
growth rate g; its weight in the stock is reported for every year.

Run alone to print the year table and the sensitivity grid:
    python3 code/facilities_stock.py
"""

from __future__ import annotations

import re
import sys
from collections import defaultdict
from pathlib import Path

import numpy as np

sys.path.insert(0, str(Path(__file__).resolve().parent))
from stage1_assumptions import (
    C30_PRESAMPLE_G,
    DELTA_PLANT,
    DELTA_SHELL,
    FACILITY_SHELL_SHARE,
)

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
C30_XLSX = RAW / "privsatime.xlsx"
NIPA_XLSX = RAW / "bea_2024" / "NIPA_Section5All_xls.xlsx"

FIRST_YEAR = 2014
REPORT_YEARS = (2018, 2022, 2023, 2024, 2025)

# Independent checks on the facility stock, $bn, from notes/us-national-accounts/dc_structures_sizing.tex.
CHECKS_BN = {
    2018: {"filings": (44, 69), "MW_x_cost": (26, 44), "BEA_514_structures": 84},
    2022: {"filings": (102, 166), "MW_x_cost": (78, 114), "BEA_514_structures": 235},
    2024: {"filings": (179, 269), "MW_x_cost": (158, 248), "BEA_514_structures": 304},
}


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def c30_annual_m(path: Path = C30_XLSX) -> dict[int, float]:
    """Annual data-center construction, $m. Mean of 12 monthly SAAR values."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    rows = list(wb["Private SA"].iter_rows(values_only=True))
    wb.close()
    header = rows[3]
    col = next(j for j, h in enumerate(header) if h and "data center" in str(h).lower())
    pat = re.compile(r"^([A-Za-z]{3})-(\d{2})[pr]?$")
    months = defaultdict(list)
    for r in rows[4:]:
        if not r or r[0] is None:
            continue
        m = pat.match(str(r[0]).strip())
        if not m:
            continue
        v = _num(r[col])
        if v is not None:
            months[2000 + int(m.group(2))].append(v)
    return {y: float(np.mean(vs)) for y, vs in sorted(months.items()) if len(vs) == 12}


def office_price_index(path: Path = NIPA_XLSX) -> dict[int, float]:
    """NIPA 5.4.4 line 4 Office (2017=100); latest year spliced from 5.3.4
    Commercial and health care when Office is not yet published."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)

    def read(sheet, pred):
        rows = list(wb[sheet].iter_rows(values_only=True))
        years = {}
        for j, v in enumerate(rows[7]):
            try:
                years[int(v)] = j
            except (TypeError, ValueError):
                pass
        for r in rows[8:]:
            line = str(r[0]).strip() if r[0] else ""
            name = str(r[1]).strip() if r[1] else ""
            if pred(line, name):
                return {y: _num(r[j]) for y, j in years.items() if j < len(r)}
        raise RuntimeError(f"price index row not found in {sheet}")

    office = read("T50404-A", lambda l, n: l == "4" and n == "Office")
    comm = read("T50304-A", lambda l, n: "Commercial and health care" in n)
    wb.close()
    p = {y: v for y, v in office.items() if v is not None and y >= FIRST_YEAR}
    last = max(p)
    y = last + 1
    while comm.get(y) is not None and comm.get(y - 1):
        p[y] = p[y - 1] * comm[y] / comm[y - 1]
        y += 1
    return p


def facility_stock(
    i_m: dict[int, float],
    prices: dict[int, float],
    *,
    shell_share: float = FACILITY_SHELL_SHARE,
    delta_shell: float = DELTA_SHELL,
    delta_plant: float = DELTA_PLANT,
    g: float | None = C30_PRESAMPLE_G,
) -> dict[int, dict[str, float]]:
    """Per year: I, K_shell, K_plant, K_fac (current $m), blended delta,
    steady-state replacement I_ss, and the share of K_fac inherited from the
    pre-sample anchor."""
    years = [y for y in sorted(i_m) if y >= FIRST_YEAR and y in prices]
    real = {y: i_m[y] * 100.0 / prices[y] for y in years}
    i0 = real[years[0]]
    if g is None:
        k_s = k_p = 0.0
    else:
        k_s = shell_share * i0 / (delta_shell + g)
        k_p = (1.0 - shell_share) * i0 / (delta_plant + g)
    a_s, a_p = k_s, k_p  # anchor components, depreciated forward
    out = {}
    for y in years:
        k_s = (1.0 - delta_shell) * k_s + shell_share * real[y]
        k_p = (1.0 - delta_plant) * k_p + (1.0 - shell_share) * real[y]
        a_s *= 1.0 - delta_shell
        a_p *= 1.0 - delta_plant
        rev = prices[y] / 100.0
        ks, kp = k_s * rev, k_p * rev
        out[y] = {
            "I": i_m[y],
            "K_shell": ks,
            "K_plant": kp,
            "K_fac": ks + kp,
            "delta": (delta_shell * ks + delta_plant * kp) / (ks + kp),
            "I_ss": delta_shell * ks + delta_plant * kp,
            "anchor_share": (a_s + a_p) * rev / (ks + kp),
            "price": prices[y],
        }
    return out


def main() -> int:
    i_m = c30_annual_m()
    p = office_price_index()
    base = facility_stock(i_m, p)
    years = [y for y in REPORT_YEARS if y in base]

    print("Data-center facility stock from C30 cumulation ($bn, current cost)")
    print(f"shell share {FACILITY_SHELL_SHARE:.2f}, delta shell {DELTA_SHELL}, plant {DELTA_PLANT}, g {C30_PRESAMPLE_G}")
    hdr = f"{'':14}" + "".join(f"{y:>9}" for y in years)
    print(hdr)
    for key in ("I", "K_shell", "K_plant", "K_fac", "I_ss"):
        print(f"{key:14}" + "".join(f"{base[y][key]/1000:9.1f}" for y in years))
    print(f"{'delta':14}" + "".join(f"{base[y]['delta']:9.3f}" for y in years))
    print(f"{'I / I_ss':14}" + "".join(f"{base[y]['I']/base[y]['I_ss']:9.2f}" for y in years))
    print(f"{'anchor share':14}" + "".join(f"{base[y]['anchor_share']:9.2f}" for y in years))

    print("\nSensitivity of K_fac to the pre-sample anchor ($bn)")
    for g in (0.05, 0.10, 0.20, 0.40, None):
        s = facility_stock(i_m, p, g=g)
        lab = "K_2013 = 0" if g is None else f"g = {g:.2f}"
        print(f"{lab:14}" + "".join(f"{s[y]['K_fac']/1000:9.1f}" for y in years))

    print("\nChecks from notes/us-national-accounts/dc_structures_sizing.tex ($bn)")
    for y, c in CHECKS_BN.items():
        if y in base:
            print(
                f"{y}: K_fac {base[y]['K_fac']/1000:.0f} | filings {c['filings'][0]}–{c['filings'][1]}"
                f" | MW×cost {c['MW_x_cost'][0]}–{c['MW_x_cost'][1]} | BEA 514 structures {c['BEA_514_structures']}"
            )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
