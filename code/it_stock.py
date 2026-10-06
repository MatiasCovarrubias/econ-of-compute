#!/usr/bin/env python3
"""Server and orchestration-software stock of the Compute node.

Source: BEA detailed fixed assets, September 2025 vintage
(data/raw/bea_2024/detailnonres_stk1.xlsx). Baseline is the hosting /
internet-publishing industry (sheet 5140), asset types EP1A mainframes,
EP1G storage, EP1H system integrators. Publishing (5110) is added only as
a Microsoft-inclusive ceiling. Orchestration software is a small ratio of
the server stock, not the industry's official software line.

Run alone:
    python3 code/it_stock.py
"""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from stage1_assumptions import (
    DELTA_IT,
    DELTA_SOFT,
    HOSTING_FA,
    NETWORK_ASSET_CODES,
    PC_ASSET_CODES,
    PUBLISH_AND_HOSTING_FA,
    SERVER_ASSET_CODES,
    SOFT_TO_IT_RATIO,
    SOFTWARE_ASSET_CODES,
)

ROOT = Path(__file__).resolve().parents[1]
STK_XLSX = ROOT / "data" / "raw" / "bea_2024" / "detailnonres_stk1.xlsx"
REPORT_YEARS = (2018, 2022, 2023, 2024)
SOFT_BAND = (0.0, 0.04, 0.10)


def _num(v):
    try:
        return float(v or 0.0)
    except (TypeError, ValueError):
        return 0.0


def load_industry_assets(path: Path, sheets: tuple[str, ...], codes: tuple[str, ...]) -> dict[str, dict[str, dict[int, float]]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    out = {}
    for sheet in sheets:
        rows = list(wb[sheet].iter_rows(values_only=True))
        years = {}
        for j, v in enumerate(rows[5]):
            try:
                years[int(v)] = j
            except (TypeError, ValueError):
                pass
        block = {c: {} for c in codes}
        for r in rows[7:]:
            if not r or r[0] is None:
                continue
            code = str(r[0]).strip()
            if code not in block:
                continue
            for y, j in years.items():
                if j < len(r):
                    block[code][y] = _num(r[j])
        out[sheet] = block
    wb.close()
    return out


def _sum(block: dict, sheets: tuple[str, ...], codes: tuple[str, ...], year: int) -> float:
    return sum(block[s][c].get(year, 0.0) for s in sheets for c in codes)


def it_stock(path: Path = STK_XLSX, years: tuple[int, ...] = REPORT_YEARS) -> dict[int, dict[str, float]]:
    codes = SERVER_ASSET_CODES + PC_ASSET_CODES + NETWORK_ASSET_CODES + SOFTWARE_ASSET_CODES
    block = load_industry_assets(path, PUBLISH_AND_HOSTING_FA, codes)
    out = {}
    for y in years:
        k_it = _sum(block, HOSTING_FA, SERVER_ASSET_CODES, y)
        k_ceil = _sum(block, PUBLISH_AND_HOSTING_FA, SERVER_ASSET_CODES, y)
        out[y] = {
            "K_IT": k_it,
            "K_IT_ceiling": k_ceil,
            "K_IT_publishing": k_ceil - k_it,
            "K_soft": SOFT_TO_IT_RATIO * k_it,
            "K_PC_hosting": _sum(block, HOSTING_FA, PC_ASSET_CODES, y),
            "K_net_hosting": _sum(block, HOSTING_FA, NETWORK_ASSET_CODES, y),
            "K_soft_hosting_official": _sum(block, HOSTING_FA, SOFTWARE_ASSET_CODES, y),
            "I_IT": DELTA_IT * k_it,
            "I_soft": DELTA_SOFT * SOFT_TO_IT_RATIO * k_it,
            "delta": (DELTA_IT * k_it + DELTA_SOFT * SOFT_TO_IT_RATIO * k_it)
            / max(k_it * (1.0 + SOFT_TO_IT_RATIO), 1e-9),
        }
    return out


def main() -> int:
    stocks = it_stock()
    years = [y for y in REPORT_YEARS if y in stocks]
    print("Compute IT stock from BEA FA, Sept 2025 vintage ($bn, current cost)")
    print(f"baseline = hosting 5140 × {SERVER_ASSET_CODES}; software = {SOFT_TO_IT_RATIO:.2f} × servers")
    hdr = f"{'':28}" + "".join(f"{y:>9}" for y in years)
    print(hdr)
    keys = (
        ("K_IT", "servers (used)"),
        ("K_IT_publishing", "publishing 5110"),
        ("K_IT_ceiling", "5110+5140 ceiling"),
        ("K_soft", f"software @ {SOFT_TO_IT_RATIO:.2f}"),
        ("I_IT", "I_ss servers"),
        ("K_PC_hosting", "5140 PCs (not used)"),
        ("K_net_hosting", "5140 networking (not used)"),
        ("K_soft_hosting_official", "5140 official software"),
    )
    for key, lab in keys:
        print(f"{lab:28}" + "".join(f"{stocks[y][key]/1000:9.1f}" for y in years))

    print("\nSoftware-ratio band on hosting servers ($bn software)")
    for r in SOFT_BAND:
        mark = " (used)" if abs(r - SOFT_TO_IT_RATIO) < 1e-9 else ""
        print(
            f"  {r:.2f}{mark:8}"
            + "".join(f"{r * stocks[y]['K_IT']/1000:9.1f}" for y in years)
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
